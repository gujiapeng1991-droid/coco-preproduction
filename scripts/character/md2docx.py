#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""轻量 Markdown -> Word（支持标题 / 表格 / 引用 / 列表 / 行内粗体），中文排版优化。"""
import re, sys
from docx import Document
from docx.shared import Pt, RGBColor, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

FONT = "PingFang SC"

def set_font(run, size=11, bold=False, color=None, name=FONT):
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.name = name
    run._element.rPr.rFonts.set(qn('w:eastAsia'), name)
    if color:
        run.font.color.rgb = RGBColor.from_string(color)

def add_runs(p, text, size=11, base_bold=False, color=None):
    """解析 **粗体** 与 `代码`"""
    for part in re.split(r'(\*\*.+?\*\*|`.+?`)', text):
        if not part:
            continue
        if part.startswith('**') and part.endswith('**'):
            set_font(p.add_run(part[2:-2]), size, True, color)
        elif part.startswith('`') and part.endswith('`'):
            set_font(p.add_run(part[1:-1]), size - 0.5, base_bold, "5B6472")
        else:
            set_font(p.add_run(part), size, base_bold, color)

def shade(cell, hexcolor):
    tcPr = cell._tc.get_or_add_tcPr()
    sh = OxmlElement('w:shd'); sh.set(qn('w:fill'), hexcolor)
    tcPr.append(sh)

def main(md_path, docx_path):
    lines = open(md_path, encoding="utf-8").read().split("\n")
    doc = Document()
    st = doc.styles["Normal"]
    st.font.name = FONT; st.font.size = Pt(11)
    st.element.rPr.rFonts.set(qn('w:eastAsia'), FONT)
    for s in doc.sections:
        s.top_margin = Cm(2.2); s.bottom_margin = Cm(2.2)
        s.left_margin = Cm(2.4); s.right_margin = Cm(2.4)

    i, n = 0, len(lines)
    while i < n:
        ln = lines[i].rstrip()
        i += 1
        if not ln.strip():
            continue

        # 表格
        if ln.strip().startswith("|") and i < n and re.match(r'^\s*\|[\s:\-\|]+\|\s*$', lines[i]):
            head = [c.strip() for c in ln.strip().strip('|').split('|')]
            i += 1
            body = []
            while i < n and lines[i].strip().startswith("|"):
                body.append([c.strip() for c in lines[i].strip().strip('|').split('|')])
                i += 1
            t = doc.add_table(rows=1, cols=len(head))
            t.style = "Table Grid"
            for j, h in enumerate(head):
                c = t.rows[0].cells[j]; c.text = ""
                add_runs(c.paragraphs[0], h, 10, True, "FFFFFF")
                shade(c, "3C4252")
            for row in body:
                cells = t.add_row().cells
                for j, v in enumerate(row[:len(head)]):
                    cells[j].text = ""
                    add_runs(cells[j].paragraphs[0], v, 10)
            doc.add_paragraph()
            continue

        # 标题
        m = re.match(r'^(#{1,4})\s+(.*)$', ln)
        if m:
            lvl, txt = len(m.group(1)), m.group(2).strip()
            p = doc.add_heading(level=min(lvl, 4))
            p.paragraph_format.space_before = Pt(14 if lvl <= 2 else 10)
            p.paragraph_format.space_after = Pt(6)
            for r in p.runs:
                r._element.getparent().remove(r._element)
            size = {1: 20, 2: 16, 3: 13.5, 4: 12}[lvl]
            col = {1: "1F2430", 2: "2B3A55", 3: "3C4252", 4: "3C4252"}[lvl]
            add_runs(p, txt, size, True, col)
            continue

        # 分隔线
        if re.match(r'^\s*---+\s*$', ln):
            doc.add_paragraph("─" * 46).runs[0].font.color.rgb = RGBColor.from_string("D5D9E0")
            continue

        # 引用
        if ln.strip().startswith(">"):
            buf = []
            while i <= n and ln.strip().startswith(">"):
                buf.append(ln.strip().lstrip('>').strip())
                ln = lines[i] if i < n else ""; i += 1
            p = doc.add_paragraph()
            p.paragraph_format.left_indent = Cm(0.6)
            p.paragraph_format.space_after = Pt(6)
            add_runs(p, " ".join(x for x in buf if x), 10.5, False, "5B6472")
            continue

        # 列表
        m = re.match(r'^\s*[-*]\s+(.*)$', ln)
        if m:
            p = doc.add_paragraph(style="List Bullet")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, m.group(1), 11)
            continue

        m = re.match(r'^\s*(\d+)\.\s+(.*)$', ln)
        if m:
            p = doc.add_paragraph(style="List Number")
            p.paragraph_format.space_after = Pt(3)
            add_runs(p, m.group(2), 11)
            continue

        p = doc.add_paragraph()
        p.paragraph_format.space_after = Pt(7)
        p.paragraph_format.line_spacing = 1.45
        add_runs(p, ln.strip(), 11)

    doc.save(docx_path)
    print("[ok]", docx_path)

if __name__ == "__main__":
    import argparse
    from pathlib import Path
    ap = argparse.ArgumentParser(
        description="Markdown -> Word（支持标题 / 表格 / 引用 / 列表 / 行内粗体与代码）")
    ap.add_argument("md", help="输入 .md 文件")
    ap.add_argument("out_pos", nargs="?", default=None,
                    help="输出 .docx 路径（位置参数写法；不给则与 md 同名）")
    ap.add_argument("-o", "--out", dest="out", default=None,
                    help="输出 .docx 路径（与位置参数等价，两种写法都支持）")
    a = ap.parse_args()
    main(a.md, a.out or a.out_pos or str(Path(a.md).with_suffix(".docx")))
